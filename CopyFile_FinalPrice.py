import tkinter as tk
from tkinter import filedialog
import openpyxl as xl
import ctypes
import os
import shutil
# Create function to mapping key-value from dict
def get_value(dictionary, string):
    for key, value in dictionary.items():
        if string in key:
            return value
    return None
root = tk.Tk()
root.withdraw()
temp_path = filedialog.askopenfilename(title="Chọn đường dẫn template Final price Excel:")
wb = xl.load_workbook(temp_path,data_only=True)
ws = wb.active
lrow = ws.max_row
tempdir = os.path.dirname(temp_path) #Get parent folder directory of Template file
dict1={} #Create a dict to mapping file - folder
for i in range(2,lrow+1):
    foldname = str(ws.cell(row=i,column=4).value.date().strftime("%d-%b")) #New folder name
    new_fol = tempdir + "/" + foldname #Directory to new folder
    if not os.path.exists(new_fol): #Check folder exist then create new folder
        os.mkdir(new_fol)
    filename = "Inv " + str(ws.cell(row=i,column=1).value) + " (final price).xlsx"
    dict1[filename] = foldname
for path, dirs, files in os.walk(tempdir): #Loop thru all folders and subfolders
    for file in files:
        if file in dict1:
            src_path = os.path.join(path, file).replace("\\","/")
            dest_path = os.path.join(tempdir, get_value(dict1,file),file).replace("\\","/")
            try:
                shutil.copyfile(src_path, dest_path) #Copy file

            # If source and destination are same
            except shutil.SameFileError:
                print("Source and destination represents the same file.")
ctypes.windll.user32.MessageBoxW(0, "Done, 520!", "Thông báo", 0)
