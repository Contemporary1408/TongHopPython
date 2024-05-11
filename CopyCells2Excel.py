import openpyxl as xl
import webview
def webview_file_dialog():
    file = None
    def open_file_dialog(w):
        nonlocal file
        try:
            file = w.create_file_dialog(webview.OPEN_DIALOG)[0]
        except TypeError:
            pass  # user exited file dialog without picking
        finally:
            w.destroy()
    window = webview.create_window("", hidden=True)
    webview.start(open_file_dialog, window)
    # file will either be a string or None
    return file
path1 = webview_file_dialog()
def webview_file_dialog1():
    file = None
    def open_file_dialog(w):
        nonlocal file
        try:
            file = w.create_file_dialog(webview.OPEN_DIALOG)[0]
        except TypeError:
            pass  # user exited file dialog without picking
        finally:
            w.destroy()
    window = webview.create_window("", hidden=True)
    webview.start(open_file_dialog, window)
    # file will either be a string or None
    return file
path2 = webview_file_dialog1()
#path1 = "C:/Users/Contemporary/Desktop/exceltest.xlsx"
#path2 = "C:/Users/Contemporary/Desktop/exceltest2.xlsx"
wb1 = xl.load_workbook(path1)
ws1 = wb1.active
wb2 = xl.load_workbook(path2)
ws2 = wb2.active
for i in range(1,18):
    ws2.cell(column=2, row=i).value = ws1.cell(column=1,row=i).value
wb2.save(path2)
