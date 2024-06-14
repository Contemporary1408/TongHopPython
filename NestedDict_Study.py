import openpyxl as xl
wb = xl.load_workbook(r"C:\Users\anh.doduc\Desktop\INV.xlsx",data_only=True)
ws = wb.active
last_row=ws.max_row
dict = {}
for i in range(2,last_row+1):
    if ws.cell(row=i,column=1).value != ws.cell(row=i-1,column=1).value:
        inv = ws.cell(row=i,column=1).value
        date = ws.cell(row=i,column=2).value.date().strftime("%d-%b")
        dict = {"inv_no":inv,"date":date}   
    item_name = ws.cell(row=i,column=3).value
    dict[item_name]={}
    dict[item_name]['item_qty'] = ws.cell(row=i,column=4).value
    dict[item_name]['item_unit_price'] = ws.cell(row=i,column=5).value
    dict[item_name]['item_ttl'] = ws.cell(row=i,column=6).value
