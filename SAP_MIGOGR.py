import openpyxl as xl
import pandas as pd
import FreeSimpleGUI as sg
import win32com.client
import sys
import subprocess
import time
from winotify import Notification
ico_path = r"\\10.118.29.7\BTMV-Data\1-ALL\05. 1611\asset\sua.ico"
toast = Notification(app_id="Tool upload MIGO GR v1.0 by Do Duc Anh",
                    title="Thông báo",
                    msg="Đã Good Receipt xong!") 
excel_path = sg.popup_get_file(title='Tool upload MIGO GR v1.0 by Do Duc Anh',
                                   message='Nhập đường dẫn file template Excel:',
                                   size = (60,150),
                                   icon=ico_path)
if excel_path is None:
    sg.popup("Bạn chưa chọn file hoặc có lỗi trên SAP",
                 title='Tool upload MIGO GR v1.0 by Do Duc Anh',
                 icon=ico_path)

#Get max row index:
df = pd.read_excel(excel_path)
column_name = 'Date Index'
max_row = len(df[column_name].dropna()) + 1
print(max_row)
#Get user SAP:
wb = xl.load_workbook(excel_path,data_only=True)
ws = wb.active
user = ws['A1'].value
pw = ws['A2'].value
server = ws['A4'].value
def MIGOGR(user,pw,server):
    try:
        path = r"C:\Program Files (x86)\SAP\FrontEnd\SAPgui\saplogon.exe"
        subprocess.Popen(path)
        time.sleep(1)

        SapGuiAuto = win32com.client.GetObject('SAPGUI')
        if not type(SapGuiAuto) == win32com.client.CDispatch:
            return

        application = SapGuiAuto.GetScriptingEngine
        if not type(application) == win32com.client.CDispatch:
            SapGuiAuto = None
            return
        connection = application.OpenConnection(server, True)

        if not type(connection) == win32com.client.CDispatch:
            application = None
            SapGuiAuto = None
            return

        session = connection.Children(0)
        if not type(session) == win32com.client.CDispatch:
            connection = None
            application = None
            SapGuiAuto = None
            return

        session.findById("wnd[0]/usr/txtRSYST-BNAME").text = user
        session.findById("wnd[0]/usr/pwdRSYST-BCODE").text = pw
        session.findById("wnd[0]").sendVKey(0)
        session.findById("wnd[0]").maximize()
        session.findById("wnd[0]/tbar[0]/okcd").text = "/nMIGO_GR"
        session.findById("wnd[0]").sendVKey(0)
        session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:0002/subSUB_FIRSTLINE:SAPLMIGO:0010/ctxtGODEFAULT_TV-BWART").text = ws.cell(row = 3,column=1).value
        #Nhập ngày tháng, số phiếu:
        for i in range(2,max_row+1):
            #i = 2
            session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:0002/subSUB_HEADER:SAPLMIGO:0101/subSUB_HEADER:SAPLMIGO:0100/tabsTS_GOHEAD/tabpOK_GOHEAD_GENERAL/ssubSUB_TS_GOHEAD_GENERAL:SAPLMIGO:0112/ctxtGOHEAD-BLDAT").text = ws.cell(row = i,column=4).value.strftime('%d.%m.%Y')
            session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:0002/subSUB_HEADER:SAPLMIGO:0101/subSUB_HEADER:SAPLMIGO:0100/tabsTS_GOHEAD/tabpOK_GOHEAD_GENERAL/ssubSUB_TS_GOHEAD_GENERAL:SAPLMIGO:0112/ctxtGOHEAD-BUDAT").text = ws.cell(row = i,column=4).value.strftime('%d.%m.%Y')
            session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:0002/subSUB_HEADER:SAPLMIGO:0101/subSUB_HEADER:SAPLMIGO:0100/tabsTS_GOHEAD/tabpOK_GOHEAD_GENERAL/ssubSUB_TS_GOHEAD_GENERAL:SAPLMIGO:0112/txtGOHEAD-BKTXT").text = ws.cell(row = i,column=5).value
            #Mở rộng bảng và input data:
            frow = ws.cell(row = i,column=6).value
            lrow = ws.cell(row = i,column=7).value
            x = 0
            session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:0002/subSUB_HEADER:SAPLMIGO:0101/btnBUTTON_HEADER_TOGGLE").press()
            session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:0004/subSUB_ITEMLIST:SAPLMIGO:0200/tblSAPLMIGOTV_GOITEM/ctxtGOITEM-MAKTX[2,0]").setFocus()
            session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:0004/subSUB_ITEMLIST:SAPLMIGO:0200/tblSAPLMIGOTV_GOITEM/ctxtGOITEM-MAKTX[2,0]").caretPosition = 0
            for q in range(frow,lrow+1):
                session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:0004/subSUB_ITEMLIST:SAPLMIGO:0200/tblSAPLMIGOTV_GOITEM/ctxtGOITEM-MAKTX[2,"+str(x)+"]").text = ws.cell(row = q,column=11).value
                session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:0004/subSUB_ITEMLIST:SAPLMIGO:0200/tblSAPLMIGOTV_GOITEM/txtGOITEM-ERFMG[4,"+str(x)+"]").text = ws.cell(row = q,column=12).value
                session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:0004/subSUB_ITEMLIST:SAPLMIGO:0200/tblSAPLMIGOTV_GOITEM/ctxtGOITEM-ERFME[5," +str(x)+"]").text = "T"
                session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:0004/subSUB_ITEMLIST:SAPLMIGO:0200/tblSAPLMIGOTV_GOITEM/ctxtGOITEM-LGOBE[7," +str(x)+"]").text = "1403"
                session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:0004/subSUB_ITEMLIST:SAPLMIGO:0200/tblSAPLMIGOTV_GOITEM/ctxtGOITEM-NAME1[6," +str(x)+"]").text = "1000"
                x = x + 1
                if x > 9:
                    session.findById("wnd[0]").sendVKey(0)
                    for m in range(0,10):
                        if session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:0004/subSUB_ITEMLIST:SAPLMIGO:0200/tblSAPLMIGOTV_GOITEM/ctxtGOITEM-MAKTX[2," +str(m)+"]").Text != None:
                            session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:0004/subSUB_ITEMLIST:SAPLMIGO:0200/tblSAPLMIGOTV_GOITEM/ctxtGOITEM-CHARG[8," +str(m)+"]").Text = "EX"
                    x = 1
                    session.findById("wnd[0]/tbar[0]/btn[82]").press()
            session.findById("wnd[0]").sendVKey(0)            
            for x in range(0,9):
                if session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:0004/subSUB_ITEMLIST:SAPLMIGO:0200/tblSAPLMIGOTV_GOITEM/ctxtGOITEM-MAKTX[2," +str(x)+"]").text != "":
                   session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:0004/subSUB_ITEMLIST:SAPLMIGO:0200/tblSAPLMIGOTV_GOITEM/ctxtGOITEM-CHARG[8," +str(x)+"]").text = "EX"
            session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:0004/subSUB_HEADER:SAPLMIGO:0102/btnOK_HEADER").press() #Collapse bang input
            session.findById("wnd[0]/tbar[1]/btn[23]").press() #post
        toast.show()
    except:
        print(sys.exc_info()[0])              
MIGOGR(user = user,pw = pw,server = server)
