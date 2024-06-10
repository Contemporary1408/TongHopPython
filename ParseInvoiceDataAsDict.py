import openpyxl

class InvoiceParser:
    def __init__(self, file_path):
        self.file_path = file_path
        self.invoices = []

    def parse(self):
        # Load the workbook and select the active sheet
        workbook = openpyxl.load_workbook(self.file_path)
        sheet = workbook.active

        # Iterate through the rows and parse the data
        for row in sheet.iter_rows(min_row=2, values_only=True):  # Assuming the first row is the header
            invoice_number, date, item, quantity, price, total = row
            self.invoices.append({
                'Invoice Number': invoice_number,
                'Date': date,
                'Item': item,
                'Quantity': int(quantity),
                'Price': float(price),
                'Total': float(total)
            })

    def get_invoices(self):
        return self.invoices

# Example usage
file_path = r"C:\Users\anh.doduc\Desktop\INV.xlsx"  # Replace with the path to your Excel file
parser = InvoiceParser(file_path)
parser.parse()
invoices = parser.get_invoices()

# Display the parsed invoices
for invoice in invoices:
    print(f"Invoice Number: {invoice['Invoice Number']}")
    print(f"Date: {invoice['Date']}")
    print(f"Item: {invoice['Item']}")
    print(f"Quantity: {invoice['Quantity']}")
    print(f"Price: {invoice['Price']}")
    print(f"Total: {invoice['Total']}")
    print("-" * 20)


# ############################################
# More specific:
import openpyxl

class InvoiceParser:
    def __init__(self, file_path):
        self.file_path = file_path
        self.invoices = []

    def parse(self):
        # Load the workbook and select the active sheet
        workbook = openpyxl.load_workbook(self.file_path)
        sheet = workbook.active

        # Get the header row
        headers = [cell.value for cell in sheet[1]]

        # Iterate through the rows and parse the data
        for row in sheet.iter_rows(min_row=2, values_only=True):  # Assuming the first row is the header
            invoice_data = {headers[i]: row[i] for i in range(len(headers))}
            self.invoices.append(invoice_data)

    def get_invoices(self):
        return self.invoices

# Example usage
file_path = 'path_to_your_excel_file.xlsx'  # Replace with the path to your Excel file
parser = InvoiceParser(file_path)
parser.parse()
invoices = parser.get_invoices()

# Each invoice is already a dictionary with all columns included
for invoice in invoices:
    print(invoice)
